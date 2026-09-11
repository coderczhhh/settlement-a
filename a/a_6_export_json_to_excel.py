from __future__ import annotations

import argparse
import json
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from a.a_utils import progress, read_json, status


FONT_NAME = "Arial Unicode MS"
TITLE_COLOR = "1F4E78"
HEADER_COLOR = "4472C4"
SECTION_COLOR = "D9EAF7"
TOTAL_COLOR = "E2F0D9"
WARNING_COLOR = "FFF2CC"
ERROR_COLOR = "F4CCCC"
TEXT_COLOR = "1F1F1F"
LIGHT_BORDER = Side(style="thin", color="D9E2F3")
NUMBER_FORMAT = "#,##0.000"


def _excel_value(value: Any) -> Any:
    """把JSON/Decimal值转换成Excel可写入的标量。"""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _style_title(cell: Any) -> None:
    cell.font = Font(name=FONT_NAME, size=14, bold=True, color=TITLE_COLOR)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def _style_section(ws: Worksheet, row: int, start_col: int, width: int, title: str) -> None:
    for column in range(start_col, start_col + width):
        cell = ws.cell(row=row, column=column)
        cell.fill = PatternFill("solid", fgColor=SECTION_COLOR)
        cell.border = Border(top=LIGHT_BORDER, bottom=LIGHT_BORDER)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=TITLE_COLOR)
        cell.alignment = Alignment(vertical="center")
    ws.cell(row=row, column=start_col, value=title)
    ws.row_dimensions[row].height = 22


def _style_headers(ws: Worksheet, row: int, start_col: int, headers: list[str]) -> None:
    for offset, header in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + offset, value=header)
        cell.fill = PatternFill("solid", fgColor=HEADER_COLOR)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=LIGHT_BORDER, right=LIGHT_BORDER)
    ws.row_dimensions[row].height = 24


def _write_market_table(
    ws: Worksheet,
    rows: list[dict[str, Any]],
    *,
    start_row: int,
    start_col: int,
    title: str,
) -> int:
    headers = ["市场主体类型", "电压等级", "求和项:售电电量", "求和项:电能电费"]
    _style_section(ws, start_row, start_col, len(headers), title)
    _style_headers(ws, start_row + 1, start_col, headers)

    for row_number, record in enumerate(rows, start=start_row + 2):
        for offset, header in enumerate(headers):
            cell = ws.cell(row=row_number, column=start_col + offset, value=_excel_value(record.get(header)))
            cell.font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
            cell.alignment = Alignment(vertical="center")
            if offset >= 2:
                cell.number_format = NUMBER_FORMAT
                cell.alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=row_number, column=start_col + 1).alignment = Alignment(
            horizontal="left", vertical="center", indent=1 if record.get("电压等级") else 0
        )

        is_summary = not record.get("电压等级") or record.get("市场主体类型") == "总计"
        if is_summary:
            for column in range(start_col, start_col + len(headers)):
                cell = ws.cell(row=row_number, column=column)
                cell.font = Font(name=FONT_NAME, size=10, bold=True, color=TEXT_COLOR)
                cell.fill = PatternFill("solid", fgColor=TOTAL_COLOR)
                cell.border = Border(top=LIGHT_BORDER)

    return start_row + 1 + len(rows)


def _write_user_table(
    ws: Worksheet,
    rows: list[dict[str, Any]],
    *,
    start_row: int,
    start_col: int,
    title: str,
) -> int:
    headers = ["用户类型", "用电量"]
    _style_section(ws, start_row, start_col, len(headers), title)
    _style_headers(ws, start_row + 1, start_col, headers)

    summary_labels = {
        "1.居民、农业电量",
        "2.工商业用户电量小计",
        "代理购电用户",
        "以代理价格收费的兜底用户",
        "市场化用户",
        "总计",
        "合计（含抽水蓄能）",
    }
    for row_number, record in enumerate(rows, start=start_row + 2):
        label = str(record.get("用户类型", ""))
        label_cell = ws.cell(row=row_number, column=start_col, value=label)
        value_cell = ws.cell(
            row=row_number,
            column=start_col + 1,
            value=_excel_value(record.get("用电量")),
        )
        label_cell.font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
        label_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            indent=1 if label.strip().startswith(("其中", "1-", "35")) else 0,
        )
        value_cell.font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
        value_cell.alignment = Alignment(horizontal="right", vertical="center")
        value_cell.number_format = NUMBER_FORMAT

        if label in summary_labels:
            for column in range(start_col, start_col + 2):
                cell = ws.cell(row=row_number, column=column)
                cell.font = Font(name=FONT_NAME, size=10, bold=True, color=TEXT_COLOR)
                cell.fill = PatternFill("solid", fgColor=TOTAL_COLOR)
                cell.border = Border(top=LIGHT_BORDER)

    return start_row + 1 + len(rows)


def _value_by_label(rows: list[dict[str, Any]], label: str) -> Any:
    for row in rows:
        if row.get("用户类型") == label:
            return row.get("用电量")
    raise KeyError(f"用户汇总中缺少标签：{label}")


def _build_summary_sheet(workbook: Workbook, final_table: dict[str, Any]) -> None:
    ws = workbook.active
    ws.title = "统计总表"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TITLE_COLOR
    ws["A2"] = "A组统计结果对比表"
    _style_title(ws["A2"])
    ws["A3"] = "单位：电量为千瓦时，电费为元；表中数值由JSON直接写入，不使用Excel公式。"
    ws["A3"].font = Font(name=FONT_NAME, size=10, italic=True, color="666666")

    sections = final_table["sections"]
    _write_market_table(
        ws,
        sections["市场主体类型与电压等级汇总"],
        start_row=5,
        start_col=1,
        title="市场主体类型与电压等级汇总",
    )
    without_end = _write_user_table(
        ws,
        sections["不含永强"],
        start_row=5,
        start_col=6,
        title="不含永强",
    )
    with_start = without_end + 2
    with_end = _write_user_table(
        ws,
        sections["含永强"],
        start_row=with_start,
        start_col=6,
        title="含永强",
    )

    status_row = with_end + 2
    ws.cell(status_row, 6, "最终勾稽状态")
    ws.cell(status_row, 7, final_table["audit"]["status"])
    for column in (6, 7):
        cell = ws.cell(status_row, column)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=TEXT_COLOR)
        cell.fill = PatternFill("solid", fgColor=TOTAL_COLOR)
        cell.border = Border(top=LIGHT_BORDER, bottom=LIGHT_BORDER)

    widths = {"A": 20, "B": 42, "C": 19, "D": 19, "E": 3, "F": 38, "G": 20}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A6:D{5 + 1 + len(sections['市场主体类型与电压等级汇总'])}"
    ws.print_title_rows = "1:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _build_market_sheet(workbook: Workbook, final_table: dict[str, Any]) -> None:
    ws = workbook.create_sheet("市场主体汇总")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "5B9BD5"
    ws["A2"] = "市场主体类型与电压等级分项汇总"
    _style_title(ws["A2"])
    rows = final_table["sections"]["市场主体类型与电压等级汇总"]
    last_row = _write_market_table(
        ws,
        rows,
        start_row=4,
        start_col=1,
        title="JSON分项明细",
    )
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:D{last_row}"


def _build_user_sheet(workbook: Workbook, final_table: dict[str, Any]) -> None:
    ws = workbook.create_sheet("用户电量汇总")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "70AD47"
    ws["A2"] = "不含永强与含永强用户电量对比"
    _style_title(ws["A2"])

    sections = final_table["sections"]
    without_rows = sections["不含永强"]
    with_rows = sections["含永强"]
    _write_user_table(ws, without_rows, start_row=4, start_col=1, title="不含永强")
    _write_user_table(ws, with_rows, start_row=4, start_col=4, title="含永强")

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 3
    ws.column_dimensions["D"].width = 38
    ws.column_dimensions["E"].width = 20

    comparison_labels = [
        "1.居民、农业电量",
        "2.工商业用户电量小计",
        "代理购电用户",
        "以代理价格收费的兜底用户",
        "市场化用户",
        "总计",
    ]
    helper_start = 23
    _style_section(ws, helper_start, 7, 3, "主要项目对比数据")
    _style_headers(ws, helper_start + 1, 7, ["项目", "不含永强", "含永强"])
    for row_number, label in enumerate(comparison_labels, start=helper_start + 2):
        ws.cell(row_number, 7, label)
        ws.cell(row_number, 8, _excel_value(_value_by_label(without_rows, label)))
        ws.cell(row_number, 9, _excel_value(_value_by_label(with_rows, label)))
        for column in range(7, 10):
            ws.cell(row_number, column).font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
        for column in (8, 9):
            ws.cell(row_number, column).number_format = NUMBER_FORMAT

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "不含永强与含永强对比"
    chart.y_axis.title = "电量（千瓦时）"
    chart.x_axis.title = "用户类型"
    chart.height = 8
    chart.width = 16
    data = Reference(ws, min_col=8, max_col=9, min_row=helper_start + 1, max_row=helper_start + 7)
    categories = Reference(ws, min_col=7, min_row=helper_start + 2, max_row=helper_start + 7)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.legend.position = "t"
    ws.add_chart(chart, "G4")
    for column, width in {"G": 34, "H": 18, "I": 18}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A6"


def _flatten_exclusions(standardized: dict[str, Any] | None) -> list[list[Any]]:
    if standardized is None:
        return []
    result: list[list[Any]] = []
    for source_group, rows in standardized.get("excluded_rows", {}).items():
        for row in rows:
            result.append(
                [
                    source_group,
                    row.get("sheet", ""),
                    row.get("来源行", ""),
                    row.get("用电分类", ""),
                    row.get("原因", row.get("reason", "")),
                ]
            )
    return result


def _build_audit_sheet(
    workbook: Workbook,
    final_table: dict[str, Any],
    standardized: dict[str, Any] | None,
) -> None:
    ws = workbook.create_sheet("核对与排除")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "A5A5A5"
    ws["A2"] = "统计结果核对与排除记录"
    _style_title(ws["A2"])

    market_audit = final_table["audit"]["market_summary"]
    user_audit = final_table["audit"]["user_volume_summary"]
    audit_rows: list[tuple[str, Any, str]] = [
        ("最终勾稽状态", final_table["audit"]["status"], "核心检查总状态"),
        ("市场汇总电量差额", market_audit.get("volume_difference"), "应为0"),
        ("市场汇总电费差额", market_audit.get("fee_difference"), "应为0"),
    ]
    for name, value in user_audit.get("checks", {}).items():
        audit_rows.append((name, value, "应为0"))
    audit_rows.extend(
        [
            ("1.5倍代理追加电量", user_audit.get("proxy_15x_added_volume"), "已计入代理购电用户"),
            ("永强有效明细合计", user_audit.get("yongqiang_detail_total"), "参与含永强汇总"),
            ("永强声明合计", user_audit.get("yongqiang_declared_total"), "来源表合计行"),
            ("永强声明与明细差额", user_audit.get("yongqiang_declared_total_difference"), "仅警告"),
            ("永强声明合计状态", user_audit.get("yongqiang_declared_total_status"), "不阻断核心流程"),
        ]
    )

    _style_section(ws, 4, 1, 3, "勾稽检查")
    _style_headers(ws, 5, 1, ["检查项", "结果", "判定说明"])
    for row_number, (name, value, note) in enumerate(audit_rows, start=6):
        ws.cell(row_number, 1, name)
        ws.cell(row_number, 2, _excel_value(value))
        ws.cell(row_number, 3, note)
        for column in range(1, 4):
            ws.cell(row_number, column).font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
            ws.cell(row_number, column).alignment = Alignment(vertical="center")
        if isinstance(value, (int, float, Decimal)):
            ws.cell(row_number, 2).number_format = NUMBER_FORMAT
        if value == "WARNING":
            ws.cell(row_number, 2).fill = PatternFill("solid", fgColor=WARNING_COLOR)
            ws.cell(row_number, 2).font = Font(name=FONT_NAME, size=10, bold=True, color="9C6500")
        elif value == "FAIL":
            ws.cell(row_number, 2).fill = PatternFill("solid", fgColor=ERROR_COLOR)
            ws.cell(row_number, 2).font = Font(name=FONT_NAME, size=10, bold=True, color="9C0006")

    exclusions = _flatten_exclusions(standardized)
    _style_section(ws, 4, 5, 5, "排除记录")
    _style_headers(ws, 5, 5, ["来源组", "工作表", "来源行", "用电分类", "排除原因"])
    if exclusions:
        for row_number, values in enumerate(exclusions, start=6):
            for offset, value in enumerate(values):
                cell = ws.cell(row_number, 5 + offset, _excel_value(value))
                cell.font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)
                cell.alignment = Alignment(vertical="center")
            if values[3] == "五、打水用电":
                for column in range(5, 10):
                    ws.cell(row_number, column).fill = PatternFill("solid", fgColor=WARNING_COLOR)
    else:
        ws.cell(6, 5, "未提供标准化JSON，无法展示排除记录")
        ws.cell(6, 5).font = Font(name=FONT_NAME, size=10, italic=True, color="666666")

    if standardized is not None:
        statistics = standardized.get("statistics", {})
        statistics_start = max(21, 7 + len(audit_rows))
        _style_section(ws, statistics_start, 1, 3, "记录数统计")
        _style_headers(ws, statistics_start + 1, 1, ["统计项", "数量/内容", "说明"])
        for row_number, (name, value) in enumerate(statistics.items(), start=statistics_start + 2):
            ws.cell(row_number, 1, name)
            ws.cell(row_number, 2, _excel_value(value))
            ws.cell(row_number, 3, "来自a_2标准化统计")
            for column in range(1, 4):
                ws.cell(row_number, column).font = Font(name=FONT_NAME, size=10, color=TEXT_COLOR)

    widths = {
        "A": 34,
        "B": 24,
        "C": 24,
        "D": 3,
        "E": 24,
        "F": 24,
        "G": 12,
        "H": 24,
        "I": 40,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A6"


def _apply_global_style(workbook: Workbook) -> None:
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != FONT_NAME:
                    cell.font = Font(
                        name=FONT_NAME,
                        size=cell.font.sz or 10,
                        bold=cell.font.bold,
                        italic=cell.font.italic,
                        color=cell.font.color,
                    )
        ws.sheet_format.defaultRowHeight = 20
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5


def a_6_export_json_to_excel(
    final_table: dict[str, Any],
    standardized: dict[str, Any] | None = None,
    output_path: str | Path = "outputs/a_group/a_6_statistics_comparison.xlsx",
) -> Path:
    """把A组最终JSON转成用于人工核对的格式化Excel工作簿。"""
    status("A6", "开始把统计JSON转换为Excel核对表")
    required_sections = {"市场主体类型与电压等级汇总", "不含永强", "含永强"}
    actual_sections = set(final_table.get("sections", {}))
    missing_sections = required_sections - actual_sections
    if missing_sections:
        raise KeyError(f"最终JSON缺少必要分区：{sorted(missing_sections)}")

    workbook = Workbook()
    workbook.properties.creator = "settlement A组自动化流程"
    workbook.properties.title = "A组统计结果核对表"
    builders = [
        ("统计总表", lambda: _build_summary_sheet(workbook, final_table)),
        ("市场主体汇总", lambda: _build_market_sheet(workbook, final_table)),
        ("用户电量汇总", lambda: _build_user_sheet(workbook, final_table)),
        ("核对与排除", lambda: _build_audit_sheet(workbook, final_table, standardized)),
    ]
    for _, build in progress(builders, desc="A6 生成Excel工作表", unit="sheet"):
        build()
    _apply_global_style(workbook)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    status("A6", f"已输出Excel核对表：{output}")
    return output


def _main() -> None:
    parser = argparse.ArgumentParser(description="把A组统计JSON转换为格式化Excel核对表")
    parser.add_argument("--final-json", required=True, help="a_5_final_table.json路径")
    parser.add_argument("--standardized-json", help="a_2_standardized_records.json路径")
    parser.add_argument(
        "--output",
        default="outputs/a_group/a_6_statistics_comparison.xlsx",
        help="输出Excel路径",
    )
    args = parser.parse_args()
    standardized = read_json(args.standardized_json) if args.standardized_json else None
    a_6_export_json_to_excel(read_json(args.final_json), standardized, args.output)


if __name__ == "__main__":
    _main()
