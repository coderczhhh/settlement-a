from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from a.a_utils import json_ready, progress, status, write_json


def _read_one_workbook(path: Path) -> dict[str, Any]:
    # 生产流程只打开一次工作簿。data_only=True 表示公式单元格只读取
    # Excel 文件中已有的计算结果，不读取、解析或保存公式文本。
    workbook = load_workbook(path, data_only=True, read_only=True)

    sheets: dict[str, Any] = {}
    for sheet in progress(
        workbook.worksheets,
        desc=f"A1 读取 {path.name}",
        unit="sheet",
    ):
        max_row = sheet.max_row
        max_column = sheet.max_column
        rows: list[dict[str, Any]] = []

        for row_number, values in enumerate(
            sheet.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_column,
                values_only=True,
            ),
            start=1,
        ):
            rows.append(
                {
                    "row_number": row_number,
                    "values": [json_ready(value) for value in values],
                }
            )

        sheets[sheet.title] = {
            "sheet_name": sheet.title,
            "state": sheet.sheet_state,
            "max_row": max_row,
            "max_column": max_column,
            "dimension": sheet.calculate_dimension(),
            "columns": [get_column_letter(i) for i in range(1, max_column + 1)],
            "rows": rows,
        }

    return {
        "file_name": path.name,
        "path": str(path.resolve()),
        "sheet_names": workbook.sheetnames,
        "sheets": sheets,
    }


def a_1_read_excel_to_dict(
    user_file: str | Path,
    yongqiang_file: str | Path,
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    """每个 Excel 只打开一次，格式化读取全部 sheet 的单元格数值。"""
    paths = [Path(user_file), Path(yongqiang_file)]
    for path in paths:
        if not path.is_file():
            raise FileNotFoundError(f"输入文件不存在：{path}")
        if path.name.startswith("~$"):
            raise ValueError(f"不能把 Excel 临时锁定文件作为输入：{path}")

    status("A1", "开始读取两个 Excel 工作簿的全部数据")
    result = {
        "schema_version": "1.1",
        "read_mode": "data_only_values_once",
        "structure": "workbook -> sheets -> rows -> values",
        "workbooks": [_read_one_workbook(path) for path in paths],
    }
    if output_path is not None:
        output = write_json(result, output_path)
        status("A1", f"已输出全量工作簿字典：{output}")
    status("A1", "全部 Excel 数据读取完成")
    return result


def _main() -> None:
    parser = argparse.ArgumentParser(description="把两个 Excel 工作簿的全部数据读取为 JSON 字典")
    parser.add_argument("--user-file", required=True)
    parser.add_argument("--yongqiang-file", required=True)
    parser.add_argument("--output", default="outputs/a_group/a_1_all_workbooks.json")
    args = parser.parse_args()
    a_1_read_excel_to_dict(args.user_file, args.yongqiang_file, args.output)


if __name__ == "__main__":
    _main()
