from __future__ import annotations

import json
import unittest
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from a.a_0_run_all import a_0_run_all


ROOT = Path(__file__).resolve().parents[2]
USER_FIXTURE = ROOT / "datafortest/2026年6月用户侧清算数据_test.xlsx"
YONGQIANG_FIXTURE = ROOT / "datafortest/永强供电公司2026年06月月度数据表_test.xlsx"


class TestAPipeline(unittest.TestCase):
    @unittest.skipUnless(
        USER_FIXTURE.is_file() and YONGQIANG_FIXTURE.is_file(),
        "公开仓库不附带业务 Excel；请在 datafortest/ 放置脱敏测试夹具",
    )
    def test_masked_workbooks(self) -> None:
        with TemporaryDirectory() as output_dir:
            result = a_0_run_all(
                USER_FIXTURE,
                YONGQIANG_FIXTURE,
                output_dir,
                Decimal("0"),
            )
            manifest = result["manifest"]
            self.assertEqual(manifest["status"], "PASS")
            self.assertNotIn("sha256", manifest["inputs"][0])
            self.assertNotIn("sha256", manifest["outputs"][0])

            raw_data = json.loads(
                (Path(output_dir) / "a_1_all_workbooks.json").read_text(encoding="utf-8")
            )
            self.assertEqual(raw_data["read_mode"], "data_only_values_once")
            self.assertNotIn("sha256", raw_data["workbooks"][0])
            first_sheet = next(iter(raw_data["workbooks"][0]["sheets"].values()))
            self.assertNotIn("formulas", first_sheet)

            standardized = json.loads(
                (Path(output_dir) / "a_2_standardized_records.json").read_text(encoding="utf-8")
            )
            self.assertEqual(
                standardized["excluded_rows"]["user_sales"][0]["用电分类"],
                "五、打水用电",
            )
            self.assertFalse(
                any(
                    record["用电分类"] == "五、打水用电"
                    for source in ("user_sales", "yongqiang")
                    for record in standardized["records"][source]
                )
            )

            self.assertEqual(manifest["record_counts"]["user_sales_rows"], 659)
            self.assertEqual(manifest["record_counts"]["user_sales_excluded_rows"], 1)
            self.assertEqual(manifest["record_counts"]["proxy_15x_rows"], 609)
            self.assertEqual(manifest["record_counts"]["yongqiang_rows"], 161)
            self.assertEqual(manifest["key_results"]["without_yongqiang_total"], Decimal("1378.394"))
            self.assertEqual(manifest["key_results"]["with_yongqiang_total"], Decimal("21181.394"))
            self.assertEqual(result["final_table"]["audit"]["status"], "PASS")

            excel_path = Path(output_dir) / "a_6_statistics_comparison.xlsx"
            self.assertEqual(result["excel_output"], excel_path)
            self.assertTrue(excel_path.is_file())
            workbook = load_workbook(excel_path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["统计总表", "市场主体汇总", "用户电量汇总", "核对与排除"],
            )
            self.assertEqual(Decimal(str(workbook["市场主体汇总"]["C43"].value)), Decimal("628.106"))
            self.assertEqual(Decimal(str(workbook["用户电量汇总"]["B22"].value)), Decimal("1378.394"))
            self.assertEqual(len(workbook["用户电量汇总"]._charts), 1)
            self.assertEqual(workbook["核对与排除"]["H6"].value, "五、打水用电")
            self.assertFalse(
                any(
                    cell.data_type == "f"
                    for worksheet in workbook.worksheets
                    for row in worksheet.iter_rows()
                    for cell in row
                )
            )


if __name__ == "__main__":
    unittest.main()
